from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
d=Document()
for s in d.sections: s.top_margin=s.bottom_margin=Cm(1.6); s.left_margin=s.right_margin=Cm(2.0)
st=d.styles['Normal']; st.font.name='Times New Roman'; st.font.size=Pt(11)
def c(t,b=True,size=13):
    p=d.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; r=p.add_run(t); r.bold=b; r.font.size=Pt(size); p.paragraph_format.space_after=Pt(2)
def para(t):
    p=d.add_paragraph(t); p.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY; p.paragraph_format.space_after=Pt(6)
def head(t):
    p=d.add_paragraph(); r=p.add_run(t); r.bold=True; p.paragraph_format.space_after=Pt(3)
def item(n,t):
    p=d.add_paragraph(f"{n}. {t}"); p.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.left_indent=Cm(0.9); p.paragraph_format.first_line_indent=Cm(-0.6); p.paragraph_format.space_after=Pt(4)
c("MEMORANDUM OF UNDERSTANDING",size=14); c("Between",False,12)
c("Sidharth Government College, Nadaun, District Hamirpur, Himachal Pradesh"); c("And",False,12)
c("Jaypee University of Information Technology, Waknaghat, Solan")

para("The Jaypee University of Information Technology (JUIT), Waknaghat was set up by Act No. 14 of 2002 vide Extraordinary Gazette notification of Government of Himachal dated May 23, 2002 and approved by the University Grants Commission under section 2(f) of the UGC Act. The University commenced academic activities from July 2002.")
para("Sidharth Government College, Nadaun is a Government degree college of the Department of Higher Education, Himachal Pradesh, located at Nadaun in District Hamirpur. The College offers undergraduate programmes to students drawn largely from the surrounding rural region and works to give them opportunities in learning, research and personal development.")
para("This MOU is set forth for collaboration between JUIT and Sidharth Government College, Nadaun for sharing scientific knowledge and expertise in areas of mutual interest, with particular emphasis on research by students and faculty. This partnership also aims at enhancing technological, social and cultural relations between both the organizations. This document presents a general framework for such a partnership which can be updated from time to time. This MOU may be modified by mutual consent of the authorized officials of the organizations. This MOU shall become effective upon signature by the authorized officials and will remain in effect until mutually modified or terminated on mutual consent by anyone of the Partners.")
head("Areas of Collaboration:")
item(1,"Joint research activities: Faculty of both organizations will jointly plan and carry out research, including research on local problems of the region, and will approach funding agencies to submit joint project proposals in the areas of Life Sciences, Biotechnology, Physics, Chemistry, Mathematics, Computer Science, Humanities or any other area of mutual interest.")
item(2,"Mentorship of student projects: Faculty of JUIT will act as co-mentors for research projects of students of Sidharth Government College. Student teams will work under a faculty mentor from each organization to frame a research question, prepare an experimental plan, collect and analyse data, and present their findings.")
item(3,"Research education and joint workshops: Both organizations will jointly organize workshops on research methodology, literature search, experimental design, data analysis, scientific paper writing and preparation of manuscripts, research ethics, plagiarism prevention and responsible use of AI tools.")
item(4,"Patents and innovation: Both organizations will identify problems of local relevance and jointly develop ideas towards filing of patents, with an initial target of two patent applications, one in Biological Sciences and one in Computer Science. Rights in jointly developed intellectual property will be shared as agreed in writing before filing.")
item(5,"Joint academic events: JUIT will host a one-day joint symposium at which students and faculty of both organizations will present posters and research ideas. Both organizations will also hold joint seminars, conferences and other academic and research events by mutual consent.")
item(6,"Follow-up and guidance: Progress of joint projects will be reviewed at intervals agreed by the coordinators, and outcomes will be presented at the joint events. JUIT will provide counselling to students and faculty on funding opportunities, research exposure and further collaboration.")
item(7,"Visits and exchanges: Both organizations will accommodate students and faculty for visits to their premises, including research facilities and laboratories, as per their policy. Faculty/students internships, staff exchanges or visiting programs will be conducted on mutually agreeable terms.")
item(8,"Any other joint academic and research activities viewed to be mutually beneficial will be conducted.")
head("Mutual Obligations")
para("The overall Memorandum does not entail any obligation for the two organizations to fund faculty/students/staff from the other organization, unless funds are available specifically for the purpose.")
head("Termination of Memorandum")
para("This MOU may be terminated with immediate effect by mutual agreement between both the partners or by one partner with two (02) months prior written notice to the other partner.")
head("Coordination")
para("In order to carry out the activities outlined in this MoU, incumbent Vice Chancellor of JUIT and Principal of Sidharth Government College, Nadaun shall appoint the Coordinators and share their information with one another. The coordinators will mutually plan and execute the activities.")

t=d.add_table(rows=1,cols=2); l,r=t.rows[0].cells
l.text=("Prof. Rajendra Kumar Sharma\nVice Chancellor\nJaypee University of Information Technology\nWaknaghat, Solan, Himachal Pradesh\nDate: ____________\n\nWitness 1:\n\nDr. Ashok Kumar Gupta\nDean (Academics & Research)\nJaypee University of Information Technology\nWaknaghat, Solan, Himachal Pradesh\n\nWitness 2:\n\nDr. Sudhir Kumar\nProfessor & Head Biotechnology & Bioinformatics\nJaypee University of Information Technology")
r.text=("Dr. Anil Kumar Thakur\nPrincipal\nSidharth Government College\nNadaun, District Hamirpur, Himachal Pradesh\nDate: ____________\n\nWitness 1:\n\nName: ______________________\nDesignation: ________________\nSidharth Government College, Nadaun\n\n\nWitness 2:\n\nName: ______________________\nDesignation: ________________\nSidharth Government College, Nadaun")
for cell in (l,r):
    for p in cell.paragraphs:
        for run in p.runs: run.font.size=Pt(10)
d.save("MOU_JUIT_Sidharth_Govt_College_Nadaun.docx")
